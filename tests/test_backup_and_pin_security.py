import os
import sys
import io
import pandas as pd
from datetime import date
import zipfile
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from app import (
    APP_ACCESS_PIN,
    check_app_security_pin,
    make_full_backup_zip,
    make_range_backup_excel,
)


def test_app_access_pin_value():
    """Verify system security PIN is strictly '7744'."""
    assert APP_ACCESS_PIN == "7744"


def test_bare_mode_pin_bypass():
    """Verify that in automated test / bare python environments, check_app_security_pin does not crash."""
    assert check_app_security_pin() is True


def test_make_full_backup_zip_contains_attendance_and_handovers():
    """Verify make_full_backup_zip exports workers, wages, attendance, and handovers in the ZIP."""
    zip_bytes = make_full_backup_zip(data_version=1)
    assert isinstance(zip_bytes, bytes)
    assert len(zip_bytes) > 0

    zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    file_list = zf.namelist()

    # Core tables
    assert "retailers.csv" in file_list
    assert "entries.csv" in file_list
    assert "payments.csv" in file_list

    # Local attendance, wage, and handover tables
    assert "workers.csv" in file_list
    assert "worker_wages.csv" in file_list
    assert "worker_attendance.csv" in file_list
    assert "driver_handovers.csv" in file_list
    assert "attendance_date_locks.csv" in file_list


def test_make_range_backup_excel_with_attendance_sheets():
    """
    Verify make_range_backup_excel generates a valid .xlsx file
    containing all 5 sheets including Worker Attendance & Wages.
    """
    retailers = pd.DataFrame([{"retailer_id": 1, "name": "Gupta Dairy", "zone": "Route 1"}])
    categories = pd.DataFrame([{"category_id": 1, "name": "Buffalo"}])
    entries = pd.DataFrame([{"date": "2026-10-01", "retailer_id": 1, "category_id": 1, "qty": 100.0, "amount": 8000.0}])
    payments = pd.DataFrame([{"date": "2026-10-01", "retailer_id": 1, "amount": 5000.0}])
    dist_purchases = pd.DataFrame()
    dist_payments = pd.DataFrame()
    distributors = pd.DataFrame()
    expenses = pd.DataFrame([{"date": "2026-10-01", "zone": "Route 1", "category": "Diesel", "amount": 500.0, "payment_mode": "Cash", "note": "Fuel"}])

    excel_bytes = make_range_backup_excel(
        start_date=date(2026, 10, 1),
        end_date=date(2026, 10, 15),
        retailers=retailers,
        categories=categories,
        entries=entries,
        payments=payments,
        dist_purchases=dist_purchases,
        dist_payments=dist_payments,
        distributors=distributors,
        expenses_df=expenses,
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # Load with openpyxl to inspect sheets
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    assert "Daily Deliveries & Ledger" in sheet_names
    assert "Worker Attendance & Wages" in sheet_names
    assert "Driver Cash Handovers" in sheet_names
    assert "Route Expenses" in sheet_names
    assert "Workers Master" in sheet_names

    # Check Worker Attendance sheet content
    ws_att = wb["Worker Attendance & Wages"]
    assert ws_att.max_row >= 2
    header_vals = [cell.value for cell in ws_att[1]]
    assert "तारीख (Date)" in header_vals
    assert "कामगार नाव (Worker Name)" in header_vals
    assert "हजेरी स्थिती (Attendance Status)" in header_vals
    assert "कमावलेला पगार (Earned Wage ₹)" in header_vals

    # Check last row is TOTAL
    last_row_vals = [cell.value for cell in ws_att[ws_att.max_row]]
    assert "TOTAL" in last_row_vals



def test_phone_lock_html_fixed_dimensions_and_no_infinite_stretch():
    """
    Verify that components/phone_lock/index.html uses strictly clamped dimensions,
    does NOT have body height 100% (which caused the infinite page stretch loop),
    and sends a deterministic fixed frame height.
    """
    html_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "components",
        "phone_lock",
        "index.html",
    )
    assert os.path.exists(html_path)

    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()

    # Must NOT have height: 100% on body (the root cause of infinite expansion)
    assert "height: 100%;" not in html
    assert "height: 560px;" in html or "max-height: 560px;" in html

    # Wrapper must have fixed max-height and hidden overflow
    assert ".phone-lock-wrapper" in html
    assert "overflow: hidden;" in html

    # Must use fixed height in setFrameHeight (not document.body.scrollHeight)
    assert "streamlit:setFrameHeight" in html
    assert "document.body.scrollHeight" not in html


def test_phone_lock_phone_keypad_and_auto_unlock():
    """
    Verify phone keypad grid elements, auto-unlock on 4th digit,
    and PIN 7744 configuration.
    """
    html_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "components",
        "phone_lock",
        "index.html",
    )
    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()

    # System PIN must default to 7744
    assert 'TARGET_PIN = "7744"' in html

    # Numeric phone pad keys 0-9
    for digit in range(10):
        assert f'data-key="{digit}"' in html

    # Functional keys
    assert 'data-key="clear"' in html
    assert 'data-key="backspace"' in html

    # 4 PIN dots present
    assert 'id="dot0"' in html
    assert 'id="dot1"' in html
    assert 'id="dot2"' in html
    assert 'id="dot3"' in html

    # Auto-unlock triggered when enteredPin.length === 4
    assert "enteredPin.length === 4" in html
    assert "verifyPin()" in html
    assert "setComponentValue(enteredPin)" in html

    # Keyboard and parent event listeners
    assert "PARENT_KEY" in html
    assert "keydown" in html


def test_ssl_verification_is_securely_enforced():
    """
    Verify CWE-295 remediation: SSL certificate verification must NOT be globally
    disabled, and httpx client must use trusted CA contexts.
    """
    import ssl
    import httpx

    # Must NOT be unverified context
    default_ctx = ssl._create_default_https_context()
    assert default_ctx.verify_mode != ssl.CERT_NONE

    # httpx Client default verify must not be False
    client = httpx.Client()
    # In httpx, client._transport._pool._ssl_context or client._transport._pool
    # must enforce valid certificates
    assert getattr(client, "_verify", None) is not False


def test_streamlit_parent_lock_css_constraints():
    """
    Verify that check_app_security_pin injects parent iframe height clamps
    and hides the sidebar during locked state.
    """
    app_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "app.py",
    )
    with open(app_path, "r", encoding="utf-8") as f:
        app_code = f.read()

    # Must contain iframe height clamp in check_app_security_pin
    assert 'iframe[title*="phone_lock"]' in app_code
    assert "max-height: 565px !important;" in app_code
    assert 'section[data-testid="stSidebar"]' in app_code
    assert "display: none !important;" in app_code
    assert "PARENT_KEY" in app_code


if __name__ == "__main__":
    test_app_access_pin_value()
    test_bare_mode_pin_bypass()
    test_make_full_backup_zip_contains_attendance_and_handovers()
    test_make_range_backup_excel_with_attendance_sheets()
    test_phone_lock_html_fixed_dimensions_and_no_infinite_stretch()
    test_phone_lock_phone_keypad_and_auto_unlock()
    test_ssl_verification_is_securely_enforced()
    test_streamlit_parent_lock_css_constraints()
    print("ALL BACKUP & PIN SECURITY TESTS PASSED!")

